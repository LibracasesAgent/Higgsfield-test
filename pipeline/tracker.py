#!/usr/bin/env python3
"""The ad tracker: one Google Sheet in Drive (Ad Creative Pipeline / "Libra Cases – Ad Tracker", data/tracker.json) with every batch
and every ad file, links to the files, and columns the team fills in (launched, spend, purchases, revenue, verdict, notes).

`make_batch.py upload` calls sync() at the end of every upload, so each Slack / Claude DM run adds its batch and ads by itself.
A sync also adds every B-numbered batch folder in Drive Outputs (Winners/ and New Products/ too) the sheet does not list yet, so a
run on a branch nobody merged still lands in the sheet.

The pipeline's Google login has the Drive scope only (the Sheets API is off in its Cloud project), so a sync rewrites the whole
file: it exports the sheet, keeps every row and every value people typed (the yellow columns, their own columns, their own tabs,
in their order), refreshes the file details of the batches it syncs, appends the new rows at the end, rebuilds the formulas and
replaces the file's content. The link stays the same. Cell colours or comments people add by hand are not kept; an edit typed in
the few seconds a sync runs can be lost.

  python3 pipeline/tracker.py show                      the sheet's link and row counts
  python3 pipeline/tracker.py sync [--batch B026]      add / refresh that batch (+ every batch in Drive the sheet lacks)
  python3 pipeline/tracker.py sync --all               refresh the file details of every batch (slow: crawls every folder)
"""
import argparse
import contextlib
import datetime
import io
import json
import os
import re
import sys
import urllib.request
import uuid

HERE = os.path.dirname(os.path.abspath(__file__))
CFG = os.path.join(HERE, "data", "tracker.json")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FOLDER = "application/vnd.google-apps.folder"

# (header, kind): d = file details (refreshed from Drive / the ledger), in = typed by people (kept), f = formula (rebuilt)
BATCH_COLS = [("Batch", "d"), ("Delivered", "d"), ("Product", "d"), ("Theme", "d"), ("Pipeline", "d"), ("What was asked", "d"),
              ("Projects", "d"), ("Statics in Drive", "f"), ("Video files in Drive", "f"), ("Drive folder", "d"), ("Notes", "d"),
              ("Status", "in"), ("Ads launched", "f"), ("Spend ($)", "f"), ("Revenue ($)", "f"), ("Purchases", "f"), ("ROAS", "f"),
              ("CPA ($)", "f"), ("Winners", "f"), ("Takeaways", "in")]
AD_COLS = [("Batch", "d"), ("Ad name", "d"), ("Type", "d"), ("Delivered", "d"), ("Product", "d"), ("Theme", "d"), ("Project", "d"),
           ("Hook", "d"), ("Angle / layout", "d"), ("Hook line / caption", "d"), ("File", "d"), ("Drive folder", "d"),
           ("Launched", "in"), ("Launch date", "in"), ("Spend ($)", "in"), ("Impressions", "in"), ("3-sec video plays", "in"),
           ("Link clicks", "in"), ("Purchases", "in"), ("Revenue ($)", "in"), ("Hook rate", "f"), ("CTR", "f"), ("CPA ($)", "f"),
           ("ROAS", "f"), ("Verdict", "in"), ("What worked", "in"), ("What didn't", "in"), ("Next action", "in")]
OURS = ("Read me", "Batches", "Ads", "Summary", "Data Sources")
EXTRA_ROWS = 400        # formula rows filled down below the data, so rows typed by hand get them too
TOKEN_PRODUCT = {"HoboBag3PieceSet": "3-Piece Set", "HoboBag": "Hobo Bag", "Hobo20": "Hobo 2.0", "VintageBag": "Vintage Bag",
                 "SilhouetteBag": "Silhouette Bag", "SlouchySoft3PieceSet": "Slouchy Set"}
KEYWORDS = [("hobo2", "Hobo 2.0"), ("set3", "3-Piece Set"), ("3piece", "3-Piece Set"), ("slouchy", "Slouchy Set"),
            ("vintage", "Vintage Bag"), ("silhouette", "Silhouette Bag"), ("hobo", "Hobo Bag"), ("remix", "Hobo Bag")]
THEMES = {"BlackFriday": "Black Friday", "CyberMonday": "Cyber Monday", "CyberWeek": "Cyber Week", "Christmas": "Christmas",
          "Gift": "Gift", "MothersDay": "Mother's Day", "Travel": "Travel", "Evergreen": "Evergreen"}
BATCH_RE = re.compile(r"^(B\d{3}(?:r\d+)?)\s*–\s*(\d{4}-\d{2}-\d{2})\s*–\s*(.+?)\s*–\s*(.+)$")


def load_cfg():
    return json.load(open(CFG, encoding="utf-8"))


def token():
    sys.path.insert(0, HERE)
    import drive_upload
    with contextlib.redirect_stdout(io.StringIO()):
        return drive_upload, drive_upload.access_token()


# ---------------------------------------------------------------- reading
def export(du, tok, sid):
    """The sheet as an openpyxl workbook (formulas kept as text)."""
    import openpyxl
    u = f"https://www.googleapis.com/drive/v3/files/{sid}/export?mimeType={XLSX}"
    with du.api("GET", u, tok) as r:
        return openpyxl.load_workbook(io.BytesIO(r.read()))


def read_tab(wb, title, cols):
    """(header, rows) of one of our tabs: rows = dicts by header, only rows holding a typed (non-formula) value."""
    if title not in wb.sheetnames:
        return [c for c, _ in cols], []
    ws = wb[title]
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    while head and not head[-1]:
        head.pop()
    fkind = {c: k for c, k in cols}
    rows = []
    for r in ws.iter_rows(min_row=2, max_col=len(head), values_only=True):
        row = {h: v for h, v in zip(head, r) if h}
        typed = [v for h, v in row.items() if fkind.get(h) != "f" and v not in (None, "") and not str(v).startswith("=")]
        if typed:
            rows.append(row)
    return head, rows


def link_of(v):
    m = re.search(r'HYPERLINK\("([^"]+)"', str(v or ""))
    return m.group(1) if m else (v if str(v or "").startswith("http") else "")


def drive_batches(du, tok, outputs):
    """Every B-numbered batch folder in Outputs, Outputs/Winners and Outputs/New Products -> {B019: {...}}."""
    found = {}
    tops = du.list_children(tok, outputs)
    subs = [(f["id"], f["name"]) for f in tops if f.get("mimeType") == FOLDER and f["name"] in ("Winners", "New Products")]
    for parent, where in [(outputs, "")] + subs:
        for f in (tops if parent == outputs else du.list_children(tok, parent)):
            m = BATCH_RE.match(f["name"]) if f.get("mimeType") == FOLDER else None
            if m:
                found.setdefault(m.group(1), dict(id=m.group(1), folder=f["id"], name=f["name"], date=m.group(2), what=m.group(3),
                                                  size=m.group(4), where=where,
                                                  link=f.get("webViewLink") or f"https://drive.google.com/drive/folders/{f['id']}"))
    return found


def crawl(du, tok, folder, skip=()):
    """Image and video files under a batch folder (recursive), newest copy per file name -> [{name, path, kind, link}]."""
    def walk(fid, path):
        out = []
        for f in du.list_children(tok, fid):
            p = f"{path}/{f['name']}" if path else f["name"]
            if f.get("mimeType") == FOLDER:
                if not any(p.startswith(s) for s in skip):
                    out += walk(f["id"], p)
            elif (f.get("mimeType") or "").split("/")[0] in ("image", "video"):
                out.append(dict(f, path=p))
        return out
    newest = {}
    for f in walk(folder, ""):
        n = os.path.splitext(f["name"])[0]
        if n not in newest or (f.get("createdTime") or "") > (newest[n].get("createdTime") or ""):
            newest[n] = f
    return [dict(name=n, path=f["path"], kind="Video" if f["mimeType"].startswith("video") else "Static",
                 link=f.get("webViewLink") or f"https://drive.google.com/file/d/{f['id']}/view")
            for n, f in sorted(newest.items(), key=lambda kv: (kv[1]["path"].count("/") > 1, kv[0]))]


# ---------------------------------------------------------------- rows from the source
def product_of(name, batch_product):
    m = re.match(r"(?:[A-Z]+_)?LC_\d{6}_([A-Za-z0-9]+?)_", name)
    if m and m.group(1) in TOKEN_PRODUCT:
        return TOKEN_PRODUCT[m.group(1)]
    if batch_product and batch_product != "Mixed":
        return batch_product
    return next((p for k, p in KEYWORDS if k in name.lower()), "Mixed")


def theme_of(name, batch_theme=""):
    m = re.match(r"LC_\d{6}_[A-Za-z0-9]+_([A-Za-z]+)_", name)
    t = m.group(1) if m else (batch_theme or "")
    return THEMES.get(t, t)


def angle_of(name, path, kind):
    proj = re.search(r"Videos/P\d+ – ([^/]+)/", path)
    if kind == "Video" and proj:
        return proj.group(1)
    m = re.search(r"_(?:S|V)\d+_(.+)$", name) or re.search(r"^(?:S|V)\d+_(.+)$", name) or re.search(r"_H\d_(.+)$", name)
    return m.group(1).replace("_", " ") if m else ""


def batch_meta(bid, led, dfolder, notes=""):
    """A Batches row (file details) from the ledger entry and / or the Drive folder name."""
    b = led.get(bid) or {}
    what, theme = (dfolder or {}).get("what", ""), ""
    for k, v in THEMES.items():               # 'Hobo Bag BlackFriday' -> product 'Hobo Bag', theme 'Black Friday'
        if what.endswith(" " + k):
            what, theme = what[: -len(k) - 1], v
    where = (dfolder or {}).get("where", "")
    pipe = b.get("pipeline") or {"Winners": "winners", "New Products": "new product"}.get(where, "catalog")
    day = b.get("date") or (dfolder or {}).get("date")
    return {"Batch": bid, "Delivered": datetime.date.fromisoformat(day) if day else None,
            "Product": b.get("product") or what or "", "Theme": theme_of("", b.get("theme")) or theme, "Pipeline": pipe,
            "What was asked": b.get("request") or (dfolder or {}).get("name") or b.get("folder") or "",
            "Projects": b.get("projects"), "Drive folder": (dfolder or {}).get("link") or b.get("link") or "", "Notes": notes}


def ad_rows(bid, meta, files, labels):
    out = []
    for f in files:
        n, kind = f["name"], f["kind"]
        pm = re.search(r"_P(\d+)_", n) or re.search(r"(?:^|_)V(\d+)_", n)
        hm = re.search(r"_H(\d)_", n)
        out.append({"Batch": bid, "Ad name": n, "Type": kind, "Delivered": meta.get("Delivered"),
                    "Product": product_of(n, meta.get("Product")), "Theme": theme_of(n, meta.get("Theme")),
                    "Project": (("P" if "_P" in n else "V") + pm.group(1)) if pm and kind == "Video" else "",
                    "Hook": f"H{hm.group(1)}" if hm else "", "Angle / layout": angle_of(n, f["path"], kind),
                    "Hook line / caption": labels.get(n, ""), "File": f["link"], "Drive folder": os.path.dirname(f["path"])})
    return out


# ---------------------------------------------------------------- writing
def build(batch_head, batches, ad_head, ads, other_tabs, synced_at):
    """The workbook: our five tabs (formulas and formats rebuilt for the columns where they now are) + people's own tabs."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    fills = {"d": PatternFill("solid", fgColor="3B2A20"), "in": PatternFill("solid", fgColor="FFE699"),
             "f": PatternFill("solid", fgColor="D9D9D9")}
    in_cell = PatternFill("solid", fgColor="FFF8DC")
    side = Side(style="thin", color="D0D0D0")
    border = Border(left=side, right=side, top=side, bottom=side)
    body, link_font = Font(name=F, size=10), Font(name=F, size=10, color="1155CC", underline="single")
    cur, pct, roas = '$#,##0.00;($#,##0.00);"-"', '0.0%;(0.0%);"-"', '0.00"x"'
    wb = Workbook()

    def head_row(ws, head, kinds):
        for i, h in enumerate(head, 1):
            c = ws.cell(row=1, column=i, value=h)
            k = kinds.get(h, "in")
            c.fill = fills[k]
            c.font = Font(name=F, bold=True, color="FFFFFF" if k == "d" else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = border
        ws.row_dimensions[1].height = 34

    # Read me
    rm = wb.active
    rm.title = "Read me"
    text = [
        ("Libra Cases ad tracker", "title"),
        ("Every batch and every ad made, with links to the files in Drive, and columns to record how each ad did on Meta.", ""),
        (f"Updated by itself at every batch upload (last update {synced_at}). New batches and ads are added at the bottom.", ""),
        ("", ""), ("Tabs", "h"),
        ("Batches: one row per batch (B001, B002 ...): date delivered, product, what was asked, the Drive folder, notes, and the results of its ads added up.", ""),
        ("Ads: one row per ad file (statics and videos), with a link to the file. Fill in the results here.", ""),
        ("Summary: results added up by product and by ad type.", ""),
        ("Data Sources: where the rows came from and when.", ""),
        ("", ""), ("How to fill it in", "h"),
        ("Yellow headers are yours to fill in (Launched, Spend, Impressions, 3-sec plays, Link clicks, Purchases, Revenue, Verdict, notes). Columns you add yourself are kept too.", "in"),
        ("Grey headers are formulas (hook rate, CTR, CPA, ROAS, and the totals on Batches and Summary). Leave them alone.", "f"),
        ("Dark headers are the batch and file details, listed from Drive. They are refreshed at each update; Notes you write are kept.", "d"),
        ("Type the Meta numbers for each ad from Ads Manager (lifetime or the period you judge). The hook rate is worked out only for videos.", ""),
        ("Cell colours and comments added by hand are not kept by the updates; use the columns instead.", ""),
        ("", ""),
        ("Example (Ads tab): Launched = Yes, Launch date = 2026-10-27, Spend = 120.50, Impressions = 15000, 3-sec plays = 4200, Link clicks = 310, Purchases = 4, Revenue = 290.00", ""),
        ("gives Hook rate 28.0%, CTR 2.1%, CPA $30.13, ROAS 2.41x.", ""),
        ("", ""), ("Verdicts", "h"),
        ("Winner: spent at least 3x the target CPA and its ROAS is above break-even. Make more of it (new hooks on the same body).", ""),
        ("Keep: still running, not enough spend to judge.", ""),
        ("Kill: spent 2-3x the target CPA with no purchases, or ROAS well under break-even.", ""),
        ("Not launched: delivered but not put on Meta.", ""),
    ]
    for i, (t, k) in enumerate(text, 1):
        c = rm.cell(row=i, column=1, value=t)
        c.font = Font(name=F, size=16 if k == "title" else 11, bold=k in ("title", "h"),
                      color="FFFFFF" if k == "d" else "000000")
        if k in fills:
            c.fill = fills[k]
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rm.column_dimensions["A"].width = 140

    def data_tab(title, head, cols, rows, key, formulas, link_cols, widths, validations, input_formats):
        ws = wb.create_sheet(title)
        kinds = dict(cols)
        head_row(ws, head, kinds)
        L = {h: get_column_letter(i) for i, h in enumerate(head, 1)}
        last = len(rows) + 1 + EXTRA_ROWS
        for r, row in enumerate(rows, 2):
            for i, h in enumerate(head, 1):
                v = row.get(h)
                if kinds.get(h) == "f":
                    continue
                if h in link_cols:
                    url = link_of(v)
                    v = f'=HYPERLINK("{url}","{link_cols[h]}")' if url else None
                c = ws.cell(row=r, column=i, value=v)
                c.font = link_font if h in link_cols else body
                c.border = border
                if kinds.get(h, "in") == "in":
                    c.fill = in_cell
                c.alignment = Alignment(wrap_text=h in ("What was asked", "Notes", "Takeaways"), vertical="top")
                if isinstance(v, (datetime.date, datetime.datetime)):
                    c.number_format = "yyyy-mm-dd"
        for r in range(2, last + 1):
            for h, (fn, fmt) in formulas.items():
                if h in L:
                    c = ws.cell(row=r, column=head.index(h) + 1, value=fn(r, L))
                    c.number_format, c.font, c.border = fmt, body, border
            if r > len(rows) + 1:
                for h in head:
                    if kinds.get(h, "in") == "in":
                        c = ws.cell(row=r, column=head.index(h) + 1)
                        c.fill, c.border, c.font = in_cell, border, body
            for h, fmt in input_formats.items():
                if h in L:
                    ws.cell(row=r, column=head.index(h) + 1).number_format = fmt
        for h, (opts, colours) in validations.items():
            if h in L:
                dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{L[h]}2:{L[h]}{last}")
                for val, col in colours.items():
                    ws.conditional_formatting.add(f"{L[h]}2:{L[h]}{last}", CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                                                    fill=PatternFill("solid", fgColor=col)))
        for h in head:
            ws.column_dimensions[L[h]].width = widths.get(h, 16)
        ws.freeze_panes = "C2" if key == "Ad name" else "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"
        return ws, L

    AL = {h: get_column_letter(i) for i, h in enumerate(ad_head, 1)}      # where each Ads column is now (people may insert columns)
    A = lambda h: "Ads!$" + AL[h] + ":$" + AL[h]          # noqa: E731

    bws, BL = data_tab(
        "Batches", batch_head, BATCH_COLS, batches, "Batch",
        {"Statics in Drive": (lambda r, L: f'=COUNTIFS({A("Batch")},${L["Batch"]}{r},{A("Type")},"Static")', "0"),
         "Video files in Drive": (lambda r, L: f'=COUNTIFS({A("Batch")},${L["Batch"]}{r},{A("Type")},"Video")', "0"),
         "Ads launched": (lambda r, L: f'=COUNTIFS({A("Batch")},${L["Batch"]}{r},{A("Launched")},"Yes")', "0"),
         "Spend ($)": (lambda r, L: f'=SUMIFS({A("Spend ($)")},{A("Batch")},${L["Batch"]}{r})', cur),
         "Revenue ($)": (lambda r, L: f'=SUMIFS({A("Revenue ($)")},{A("Batch")},${L["Batch"]}{r})', cur),
         "Purchases": (lambda r, L: f'=SUMIFS({A("Purchases")},{A("Batch")},${L["Batch"]}{r})', "0"),
         "ROAS": (lambda r, L: f'=IF({L["Spend ($)"]}{r}=0,"",{L["Revenue ($)"]}{r}/{L["Spend ($)"]}{r})', roas),
         "CPA ($)": (lambda r, L: f'=IF({L["Purchases"]}{r}=0,"",{L["Spend ($)"]}{r}/{L["Purchases"]}{r})', cur),
         "Winners": (lambda r, L: f'=COUNTIFS({A("Batch")},${L["Batch"]}{r},{A("Verdict")},"Winner")', "0")},
        {"Drive folder": "Open folder"},
        {"Batch": 8, "Delivered": 11, "Product": 14, "Theme": 12, "Pipeline": 11, "What was asked": 48, "Projects": 9, "Drive folder": 12,
         "Notes": 52, "Status": 13, "Takeaways": 40},
        {"Status": (["Delivered", "Launched", "Partly launched", "Not used"], {})}, {})
    bws["A1"].comment = Comment(f"Source: pipeline/data/batch_ledger.json and the Google Drive Outputs batch folders, updated {synced_at} "
                                f"by the pipeline (tracker.py).", "Claude")
    aws, _ = data_tab(
        "Ads", ad_head, AD_COLS, ads, "Ad name",
        {"Hook rate": (lambda r, L: f'=IF(OR({L["Type"]}{r}<>"Video",{L["Impressions"]}{r}=0),"",{L["3-sec video plays"]}{r}/{L["Impressions"]}{r})', pct),
         "CTR": (lambda r, L: f'=IF({L["Impressions"]}{r}=0,"",{L["Link clicks"]}{r}/{L["Impressions"]}{r})', pct),
         "CPA ($)": (lambda r, L: f'=IF({L["Purchases"]}{r}=0,"",{L["Spend ($)"]}{r}/{L["Purchases"]}{r})', cur),
         "ROAS": (lambda r, L: f'=IF({L["Spend ($)"]}{r}=0,"",{L["Revenue ($)"]}{r}/{L["Spend ($)"]}{r})', roas)},
        {"File": "Open"},
        {"Batch": 8, "Ad name": 52, "Type": 8, "Delivered": 11, "Product": 13, "Theme": 12, "Project": 8, "Hook": 6,
         "Angle / layout": 18, "Hook line / caption": 44, "File": 7, "Drive folder": 26, "Launched": 10, "Launch date": 11,
         "Spend ($)": 11, "Impressions": 12, "3-sec video plays": 12, "Link clicks": 10, "Purchases": 10, "Revenue ($)": 11,
         "Hook rate": 10, "CTR": 8, "CPA ($)": 10, "ROAS": 8, "Verdict": 12, "What worked": 30, "What didn't": 30, "Next action": 30},
        {"Launched": (["Yes", "No"], {}),
         "Verdict": (["Winner", "Keep", "Kill", "Not launched"], {"Winner": "C6EFCE", "Kill": "F4CCCC", "Keep": "FCE8B2"})},
        {"Launch date": "yyyy-mm-dd", "Spend ($)": cur, "Revenue ($)": cur, "Impressions": "#,##0", "3-sec video plays": "#,##0",
         "Link clicks": "#,##0", "Purchases": "#,##0"})
    aws["A1"].comment = Comment(f"Source: Google Drive, Ad Creative Pipeline / Outputs batch folders (one row per image / video file, the "
                                f"newest copy of a name), hook lines and captions from pipeline/data/batch_ledger.json; updated {synced_at} by the "
                                f"pipeline (tracker.py).", "Claude")

    # Summary
    sm = wb.create_sheet("Summary")
    prods = list(dict.fromkeys(["Hobo Bag", "Hobo 2.0", "Vintage Bag", "3-Piece Set", "Slouchy Set", "Silhouette Bag"]
                               + sorted({str(x.get("Product") or "") for x in ads} - {""})))
    cols_s = ["Ads delivered", "Ads launched", "Spend ($)", "Revenue ($)", "Purchases", "ROAS", "CPA ($)", "Winners"]

    def table(top, key, keys, title, first):
        sm.cell(row=top, column=1, value=title).font = Font(name=F, bold=True, size=12)
        for i, c in enumerate([first] + cols_s, 1):
            cell = sm.cell(row=top + 1, column=i, value=c)
            cell.fill = fills["d"] if i == 1 else fills["f"]
            cell.font = Font(name=F, bold=True, color="FFFFFF" if i == 1 else "000000")
            cell.border = border
        for j, k in enumerate(keys):
            r = top + 2 + j
            vals = [k, f"=COUNTIF({A(key)},$A{r})", f'=COUNTIFS({A(key)},$A{r},{A("Launched")},"Yes")',
                    f'=SUMIFS({A("Spend ($)")},{A(key)},$A{r})', f'=SUMIFS({A("Revenue ($)")},{A(key)},$A{r})',
                    f'=SUMIFS({A("Purchases")},{A(key)},$A{r})', f'=IF(D{r}=0,"",E{r}/D{r})', f'=IF(F{r}=0,"",D{r}/F{r})',
                    f'=COUNTIFS({A(key)},$A{r},{A("Verdict")},"Winner")']
            for c, v in enumerate(vals, 1):
                cell = sm.cell(row=r, column=c, value=v)
                cell.font, cell.border = body, border
            for c in (4, 5, 8):
                sm.cell(row=r, column=c).number_format = cur
            sm.cell(row=r, column=7).number_format = roas
        return top + 2 + len(keys)

    end = table(1, "Product", prods, "Results by product", "Product")
    table(end + 2, "Type", ["Static", "Video"], "Results by type", "Type")
    for col, w in zip("ABCDEFGHI", (16, 13, 13, 12, 12, 11, 9, 10, 9)):
        sm.column_dimensions[col].width = w

    # Data Sources
    ds = wb.create_sheet("Data Sources")
    dh = ["Written to", "Source", "Object", "Obtained via", "Query or request", "Run at", "Row count", "Changes after retrieval", "Notes"]
    for i, h in enumerate(dh, 1):
        c = ds.cell(row=1, column=i, value=h)
        c.fill, c.font, c.border = fills["d"], Font(name=F, bold=True, color="FFFFFF"), border
    rows_ds = [[f"Ads (file details: {', '.join(h for h, k in AD_COLS if k == 'd')})", "Google Drive (agent@libracases.com)",
                "Ad Creative Pipeline / Outputs batch folders, recursively", "Drive API, by the pipeline (tracker.py) at each upload",
                "Every image and video file in each batch folder, with its link", synced_at, len(ads),
                "Image and video files only; a name uploaded twice in a batch keeps its newest copy; product and theme read from the file name",
                "Hook lines and captions from pipeline/data/batch_ledger.json"],
               [f"Batches (file details)", "Git repo LibracasesAgent/Higgsfield-test + Google Drive", "pipeline/data/batch_ledger.json; the "
                "Outputs batch folder names", "Read by the pipeline (tracker.py)", "Every batch made", synced_at, len(batches),
                "A batch missing from the ledger is read from its folder name", "The first 17 batches were listed by Claude on 2026-10-25"]]
    for r, v in enumerate(rows_ds, 2):
        for c, x in enumerate(v, 1):
            cell = ds.cell(row=r, column=c, value=x)
            cell.font, cell.border = body, border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate((30, 26, 34, 26, 30, 18, 10, 40, 34), 1):
        ds.column_dimensions[get_column_letter(i)].width = w

    for title, cells in other_tabs:        # people's own tabs: their values and formulas, as they were
        ws = wb.create_sheet(title)
        for (r, c), v in cells.items():
            ws.cell(row=r, column=c, value=v)
    wb._sheets = [wb[n] for n in OURS if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in OURS]   # our tabs first
    return wb


def replace(du, tok, sid, wb):
    buf = io.BytesIO()
    wb.save(buf)
    b = uuid.uuid4().hex
    body = (f"--{b}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{{}}\r\n--{b}\r\nContent-Type: {XLSX}\r\n\r\n").encode() \
        + buf.getvalue() + f"\r\n--{b}--\r\n".encode()
    req = urllib.request.Request(f"https://www.googleapis.com/upload/drive/v3/files/{sid}?uploadType=multipart&supportsAllDrives=true"
                                 f"&fields=id,modifiedTime", data=body, method="PATCH",
                                 headers={"Authorization": f"Bearer {tok}", "Content-Type": f"multipart/related; boundary={b}"})
    with urllib.request.urlopen(req, timeout=600) as r:
        return json.loads(r.read())


# ---------------------------------------------------------------- sync
def sync(batches=(), refresh_all=False, notes=None, dry=False, quiet=False):
    """Add / refresh the given batch ids (and every batch folder in Drive the sheet lacks) -> summary dict."""
    say = (lambda *x: None) if quiet else (lambda *x: print(*x, file=sys.stderr))
    cfg = load_cfg()
    du, tok = token()
    sid = cfg["sheet_id"]
    outputs = du.load_config()["drive"]["outputs_folder"]
    wb = export(du, tok, sid)
    bhead, brows = read_tab(wb, "Batches", BATCH_COLS)
    ahead, arows = read_tab(wb, "Ads", AD_COLS)
    other = [(ws.title, {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None})
             for ws in wb.worksheets if ws.title not in OURS]
    for h, _ in BATCH_COLS:
        bhead += [h] if h not in bhead else []
    for h, _ in AD_COLS:
        ahead += [h] if h not in ahead else []

    sys.path.insert(0, HERE)
    import ledger
    led = {ledger.bid(b["n"], b.get("rev")): b for b in ledger.load()["batches"]}
    labels = {a["name"]: a.get("label") or "" for b in led.values() for a in b.get("ads") or []}
    dfold = drive_batches(du, tok, outputs)
    have = {str(r.get("Batch")) for r in brows}
    todo = {b for b in batches if b} | {b for b in dfold if b not in have}
    if refresh_all:
        todo |= have | set(dfold)
    say(f"tracker: syncing {', '.join(sorted(todo)) or 'nothing new'}")
    bidx = {str(r.get("Batch")): r for r in brows}
    aidx = {str(r.get("Ad name")): r for r in arows}
    dkind = dict(BATCH_COLS)
    akind = dict(AD_COLS)
    added_b, added_a, refreshed = [], 0, 0
    overrides = cfg.get("folder_overrides") or {}
    for bid in sorted(todo, key=lambda x: (int(re.sub(r"\D", "", x[:4]) or 0), x)):
        df = dfold.get(bid)
        folder = overrides.get(bid) or (df or {}).get("folder")
        if not folder:
            m = re.search(r"folders/([A-Za-z0-9_-]+)", (led.get(bid) or {}).get("link") or "")
            folder = m.group(1) if m else None
        if not folder:
            say(f"tracker: {bid} has no Drive folder: skipped")
            continue
        meta = batch_meta(bid, led, df, (notes or {}).get(bid, ""))
        if overrides.get(bid):
            meta["Drive folder"] = f"https://drive.google.com/drive/folders/{folder}"
        old = bidx.get(bid)
        if old is None:
            row = {h: meta.get(h) for h in bhead}
            brows.append(row)
            bidx[bid] = row
            added_b.append(bid)
        else:
            for h, v in meta.items():          # file details refreshed; Notes only filled when empty (people edit them)
                if dkind.get(h) == "d" and v not in (None, "") and not (h == "Notes" and old.get(h)):
                    old[h] = v
        files = crawl(du, tok, folder, skip=tuple(cfg.get("skip_paths", {}).get(bid, [])))
        for x in ad_rows(bid, meta, files, labels):
            o = aidx.get(x["Ad name"])
            if o is None:
                arows.append(x)
                aidx[x["Ad name"]] = x
                added_a += 1
            else:
                for h, v in x.items():
                    if akind.get(h) == "d" and v not in (None, ""):
                        o[h] = v
                refreshed += 1
    for r in arows:                            # hook lines the ledger knows now
        if not r.get("Hook line / caption") and labels.get(str(r.get("Ad name"))):
            r["Hook line / caption"] = labels[str(r.get("Ad name"))]
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    res = dict(link=cfg.get("link"), batches_added=added_b, ads_added=added_a, ads_refreshed=refreshed,
               batches=len(brows), ads=len(arows))
    if dry:
        return dict(res, dry_run=True)
    out = build(bhead, brows, ahead, arows, other, now)
    replace(du, tok, sid, out)
    chk = export(du, tok, sid)                 # verify: every row is there after the rewrite
    _, b2 = read_tab(chk, "Batches", BATCH_COLS)
    _, a2 = read_tab(chk, "Ads", AD_COLS)
    res.update(verified=len(b2) >= len(brows) and len(a2) >= len(arows))
    say(f"tracker: {len(added_b)} batch(es) and {added_a} ad(s) added, {refreshed} refreshed: {cfg.get('link')}")
    return res


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sync")
    s.add_argument("--batch", action="append", default=[], help="a batch id to add / refresh (repeatable)")
    s.add_argument("--all", action="store_true", help="refresh the file details of every batch")
    s.add_argument("--dry-run", action="store_true")
    sub.add_parser("show")
    a = ap.parse_args()
    if a.cmd == "show":
        cfg = load_cfg()
        du, tok = token()
        wb = export(du, tok, cfg["sheet_id"])
        print(json.dumps(dict(link=cfg["link"], batches=len(read_tab(wb, "Batches", BATCH_COLS)[1]),
                              ads=len(read_tab(wb, "Ads", AD_COLS)[1]), tabs=wb.sheetnames), indent=1))
        return 0
    print(json.dumps(sync([x.upper() for x in a.batch], refresh_all=a.all, dry=a.dry_run), indent=1, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
